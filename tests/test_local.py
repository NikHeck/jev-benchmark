import argparse
import json
from pathlib import Path

from openpyxl import Workbook

from build_category_input import extract_categories
from benchmark import (
    Category,
    ClassificationError,
    DEEPSEEK_OFFPEAK_PRICES,
    DEEPSEEK_PEAK_PRICES,
    MULTILINGUAL_NOTE,
    Stats,
    Usage,
    build_output,
    deepseek_cost,
    known_cost_summary,
    load_categories,
    load_tests,
    PROVIDERS,
    OPENAI_PRICES,
    RESPONSES_REASONING_EFFORT,
    openai_cost,
)


def sample_categories() -> list[Category]:
    return [
        Category(0, "01", "Food", None, 1, (1, 3), False),
        Category(1, "01.1", "Food products", 0, 2, (2,), False),
        Category(2, "01.1.1", "Bread", 1, 3, (), True),
        Category(3, "01.2", "Beverages", 0, 2, (), True),
    ]


def test_extract_categories_builds_tree(tmp_path: Path) -> None:
    path = tmp_path / "coicop.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["COICOP code", "Title", "Other"])
    sheet.append(["01", "Food and non-alcoholic beverages", "x"])
    sheet.append(["01.1", "Food", "x"])
    sheet.append(["01.1.1", "Cereals and cereal products", "x"])
    sheet.append(["01.2", "Non-alcoholic beverages", "x"])
    workbook.save(path)

    categories = extract_categories(path)
    assert categories == [
        {
            "id": 0,
            "code": "01",
            "title": "Food and non-alcoholic beverages",
            "parent_id": None,
            "level": 1,
            "is_optional_detail": False,
            "children_ids": [1, 3],
            "is_leaf": False,
        },
        {
            "id": 1,
            "code": "01.1",
            "title": "Food",
            "parent_id": 0,
            "level": 2,
            "is_optional_detail": False,
            "children_ids": [2],
            "is_leaf": False,
        },
        {
            "id": 2,
            "code": "01.1.1",
            "title": "Cereals and cereal products",
            "parent_id": 1,
            "level": 3,
            "is_optional_detail": False,
            "children_ids": [],
            "is_leaf": True,
        },
        {
            "id": 3,
            "code": "01.2",
            "title": "Non-alcoholic beverages",
            "parent_id": 0,
            "level": 2,
            "is_optional_detail": False,
            "children_ids": [],
            "is_leaf": True,
        },
    ]


def test_extract_categories_reads_selected_sheet_and_columns(tmp_path: Path) -> None:
    path = tmp_path / "custom.xlsx"
    workbook = Workbook()
    workbook.active.append(["code", "title"])
    workbook.active.append(["02", "Ignored category"])
    sheet = workbook.create_sheet("Categories")
    sheet.append(["title", "notes", "code"])
    sheet.append(["Food", "ignored", "01"])
    sheet.append(["Bread", "ignored", "01.1"])
    workbook.save(path)

    categories = extract_categories(path, sheet_number=2, code_column=3, title_column=1)
    assert [(c["code"], c["title"]) for c in categories] == [("01", "Food"), ("01.1", "Bread")]
    assert categories[1]["parent_id"] == 0


def test_extract_categories_collapses_zero_only_duplicate_level(tmp_path: Path) -> None:
    path = tmp_path / "coicop_zero.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["COICOP code", "Title"])
    sheet.append(["01", "Food and non-alcoholic beverages"])
    sheet.append(["01.2", "Non-alcoholic beverages"])
    sheet.append(["01.2.3", "Tea, maté and other plant-derived products for infusion"])
    sheet.append(["01.2.3.0", "Tea, maté and other plant-derived products for infusion"])
    workbook.save(path)

    categories = extract_categories(path)
    assert [item["code"] for item in categories] == ["01", "01.2", "01.2.3.0"]
    assert [item["id"] for item in categories] == [0, 1, 2]
    assert categories[2]["collapsed_codes"] == ["01.2.3"]
    assert categories[2]["parent_id"] == 1
    assert categories[1]["children_ids"] == [2]
    assert categories[2]["level"] == 4
    assert categories[2]["is_optional_detail"] is False
    assert categories[2]["is_leaf"] is True


def test_extract_categories_does_not_collapse_different_titles(tmp_path: Path) -> None:
    path = tmp_path / "coicop_nonduplicate.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["COICOP code", "Title"])
    sheet.append(["01", "Food"])
    sheet.append(["01.1", "Parent title"])
    sheet.append(["01.1.0", "Different child title"])
    workbook.save(path)

    categories = extract_categories(path)
    assert [item["code"] for item in categories] == ["01", "01.1", "01.1.0"]


def test_extract_categories_collapses_semantic_duplicate_with_durability_marker(tmp_path: Path) -> None:
    path = tmp_path / "coicop_semantic_duplicate.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["COICOP code", "Title"])
    sheet.append(["01", "Food and non-alcoholic beverages"])
    sheet.append(["01.1", "Food"])
    sheet.append(["01.1.1", "Cereals and cereal products (ND)"])
    sheet.append(["01.1.1.4", "Breakfast cereals (ND)"])
    sheet.append(["01.1.1.4.0", "Breakfast cereals"])
    workbook.save(path)

    categories = extract_categories(path)
    by_code = {item["code"]: item for item in categories}
    assert "01.1.1.4" not in by_code
    assert by_code["01.1.1.4.0"]["collapsed_codes"] == ["01.1.1.4"]
    assert by_code["01.1.1.4.0"]["is_optional_detail"] is True


def test_extract_categories_collapses_duplicate_chain_transitively(tmp_path: Path) -> None:
    path = tmp_path / "coicop_chain.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["COICOP code", "Title"])
    sheet.append(["01", "Food and non-alcoholic beverages"])
    sheet.append(["01.3", "Processing services"])
    sheet.append(["01.3.0", "Processing services (S)"])
    sheet.append(["01.3.0.0", "Processing services"])
    workbook.save(path)

    categories = extract_categories(path)
    assert [item["code"] for item in categories] == ["01", "01.3.0.0"]
    assert categories[1]["collapsed_codes"] == ["01.3", "01.3.0"]
    assert categories[1]["parent_id"] == 0


def test_extract_categories_defaults_to_household_divisions_01_to_13(tmp_path: Path) -> None:
    path = tmp_path / "coicop_scope.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["COICOP code", "Title"])
    sheet.append(["01", "Household food"])
    sheet.append(["13", "Other household expenditure"])
    sheet.append(["14", "NPISH expenditure"])
    sheet.append(["15", "Government expenditure"])
    workbook.save(path)

    household = extract_categories(path)
    assert [item["code"] for item in household] == ["01", "13"]

    all_divisions = extract_categories(path, household_only=False)
    assert [item["code"] for item in all_divisions] == ["01", "13", "14", "15"]


def test_load_categories_reads_explicit_tree_metadata(tmp_path: Path) -> None:
    path = tmp_path / "categories.json"
    path.write_text(
        json.dumps(
            [
                {
                    "id": 0, "code": "01", "title": "Food",
                    "parent_id": None, "level": 1, "children_ids": [1], "is_leaf": False,
                },
                {
                    "id": 1, "code": "01.1.0", "title": "Food products",
                    "parent_id": 0, "level": 3, "children_ids": [], "is_leaf": True,
                },
            ]
        ),
        encoding="utf-8",
    )
    categories = load_categories(path)
    assert categories[0].children_ids == (1,)
    assert categories[1].parent_id == 0
    assert categories[1].level == 3
    assert categories[1].is_leaf is True


def test_success_rate_counts_wrong_predictions_and_api_errors() -> None:
    stats = Stats()
    stats.record_prediction(0.1, Usage(input_tokens=10), correct=False)
    stats.record_prediction(0.1, Usage(input_tokens=10), correct=True)
    stats.record_error(0.1, RuntimeError("timeout"))

    assert stats.wrong_predictions == 1
    output = stats.to_output()
    assert output["success_rate"] == 1 / 3
    assert "hierarchy_accuracy" not in output


def test_deepseek_offpeak_is_half_peak() -> None:
    usage = Usage(cache_hit_tokens=1000, cache_miss_tokens=2000, output_tokens=300)
    peak = deepseek_cost(usage, DEEPSEEK_PEAK_PRICES)
    offpeak = deepseek_cost(usage, DEEPSEEK_OFFPEAK_PRICES)
    assert offpeak == peak / 2


def test_failure_cost_accounting_distinguishes_known_and_unknown() -> None:
    stats = Stats()

    stats.record_prediction(
        0.1, Usage(input_tokens=100, output_tokens=10), correct=True
    )
    stats.record_prediction(
        0.2, Usage(input_tokens=110, output_tokens=11), correct=False
    )
    stats.record_error(
        0.3,
        ClassificationError(
            "bad JSON", usage=Usage(input_tokens=120, output_tokens=12), cost_complete=True
        ),
    )
    stats.record_error(0.4, RuntimeError("timeout"))

    assert stats.samples == 4
    assert stats.successes == 1
    assert stats.wrong_predictions == 1
    assert stats.api_errors == 2
    assert stats.known_cost_attempts == 3
    assert stats.unknown_cost_attempts == 1
    assert stats.usage.input_tokens == 330
    assert stats.complete_cost_usage.input_tokens == 330


def test_known_cost_summary_marks_total_as_lower_bound_when_needed() -> None:
    stats = Stats(known_cost_attempts=3, unknown_cost_attempts=1)
    stats.successes = 1
    stats.wrong_predictions = 1
    stats.api_errors = 2
    summary = known_cost_summary(0.12, 0.09, stats)
    descriptions = summary.pop("descriptions")
    assert summary == {
        "known_total_usd": 0.12,
        "complete_cost_attempts_total_usd": 0.09,
        "known_cost_per_known_attempt_usd": 0.03,
        "known_cost_lower_bound_per_attempt_usd": 0.03,
        "unknown_cost_attempts": 1,
    }
    assert set(descriptions) == set(summary)
    assert all(
        isinstance(value, str) and value.strip()
        for value in descriptions.values()
    )


def test_partial_usage_from_unknown_cost_error_not_in_complete_average() -> None:
    stats = Stats()
    stats.record_prediction(
        0.1, Usage(input_tokens=100, output_tokens=10), correct=True
    )
    stats.record_error(
        0.2,
        ClassificationError(
            "later hierarchical request timed out",
            usage=Usage(input_tokens=40, output_tokens=4),
            cost_complete=False,
        ),
    )
    assert stats.known_cost_attempts == 1
    assert stats.unknown_cost_attempts == 1
    assert stats.usage.input_tokens == 140
    assert stats.complete_cost_usage.input_tokens == 100


def test_multilingual_note_is_explicit() -> None:
    assert "any language" in MULTILINGUAL_NOTE
    assert "not necessarily English" in MULTILINGUAL_NOTE


def test_load_tests_from_csv_supports_commas_utf8_and_ignores_extra_columns(tmp_path: Path) -> None:
    path = tmp_path / "input.csv"
    path.write_text(
        'sample_id,category_id,title,notes\n99,2,"Bio Vollmilch 3,8% 1L",keep\n100,3,Crème de soin,ignore me\n',
        encoding="utf-8",
    )
    tests = load_tests(path, {0, 1, 2, 3})
    assert [(test.category_id, test.title) for test in tests] == [
        (2, "Bio Vollmilch 3,8% 1L"),
        (3, "Crème de soin"),
    ]


def test_load_tests_from_csv_header_is_whitespace_and_case_tolerant(tmp_path: Path) -> None:
    path = tmp_path / "input.csv"
    path.write_text(' CATEGORY_ID , Title \n2,Bread\n', encoding="utf-8")
    tests = load_tests(path, {2})
    assert tests[0].category_id == 2
    assert tests[0].title == "Bread"


def test_load_tests_from_csv_rejects_missing_required_columns(tmp_path: Path) -> None:
    path = tmp_path / "input.csv"
    path.write_text("id,title\n2,Bread\n", encoding="utf-8")
    try:
        load_tests(path, {2})
    except ValueError as exc:
        assert "category_id" in str(exc)
        assert "missing required column" in str(exc)
    else:
        raise AssertionError("Expected ValueError")


def test_number_samples_means_full_dataset_iterations(tmp_path: Path) -> None:
    args = argparse.Namespace(
        number_samples=5,
        input_file=tmp_path / "input.csv",
        output_file=tmp_path / "output.json",
        category_file=tmp_path / "category_input.json",
    )
    stats = {
        "openai_luna.direct": Stats(),
        "openai_luna.recursive": Stats(),
        "deepseek_flash.direct": Stats(),
        "deepseek_flash.recursive": Stats(),
        "typesafe_jev.recursive": Stats(),
    }
    output = build_output(stats, args, test_count=30)
    config = output["configuration"]
    assert config["number_samples"] == 5
    assert config["full_dataset_iterations"] == 5
    assert config["tests_in_input_file"] == 30
    assert config["classifications_per_strategy"] == 150
    assert config["request_timeout_seconds"] == 600.0
    assert config["classification_target"] == "leaf_category"
    assert "complete dataset iterations" in config["sampling"]
    for provider in ("openai_luna", "deepseek_flash"):
        assert output["models"][provider]["direct"]["reasoning_effort"] == (
            RESPONSES_REASONING_EFFORT
        )


def test_build_output_prediction_log_is_separate_top_level_element(tmp_path: Path) -> None:
    args = argparse.Namespace(
        number_samples=1,
        input_file=tmp_path / "input.csv",
        output_file=tmp_path / "output.json",
        category_file=tmp_path / "category_input.json",
    )
    stats = {
        "openai_luna.direct": Stats(),
        "openai_luna.recursive": Stats(),
        "deepseek_flash.direct": Stats(),
        "deepseek_flash.recursive": Stats(),
        "typesafe_jev.recursive": Stats(),
    }
    log = [
        {
            "iteration": 1,
            "sample_index": 1,
            "model": "openai_luna",
            "strategy": "direct",
            "title": "Bread",
            "expected": {"category_id": 2, "code": "01.1.1", "title": "Bread"},
            "predicted": {"category_id": 2, "code": "01.1.1", "title": "Bread"},
            "correct": True,
            "status": "correct",
            "error": None,
        }
    ]
    output = build_output(stats, args, test_count=1, prediction_log=log)
    assert output["prediction_log"] == log
    assert "prediction_log" not in output["models"]


def test_build_output_defaults_prediction_log_to_empty_list(tmp_path: Path) -> None:
    args = argparse.Namespace(
        number_samples=1,
        input_file=tmp_path / "input.csv",
        output_file=tmp_path / "output.json",
        category_file=tmp_path / "category_input.json",
    )
    stats = {
        "openai_luna.direct": Stats(),
        "openai_luna.recursive": Stats(),
        "deepseek_flash.direct": Stats(),
        "deepseek_flash.recursive": Stats(),
        "typesafe_jev.recursive": Stats(),
    }
    output = build_output(stats, args, test_count=1)
    assert output["prediction_log"] == []


def test_selected_sol_strategy_has_own_model_effort_and_price(tmp_path: Path) -> None:
    args = argparse.Namespace(
        number_samples=1,
        input_file=tmp_path / "input.csv",
        category_file=tmp_path / "category_input.json",
    )
    output = build_output({"openai_sol.recursive": Stats()}, args, test_count=1)
    assert list(output["models"]) == ["openai_sol"]
    sol = output["models"]["openai_sol"]["recursive"]
    assert sol["model"] == PROVIDERS["openai_sol"].model == "gpt-6.1-sol"
    assert sol["reasoning_effort"] == PROVIDERS["openai_sol"].reasoning_effort == "low"
    assert sol["cost"]["prices_usd_per_1m_tokens"] == OPENAI_PRICES["openai_sol"]
    assert openai_cost(Usage(input_tokens=1_000_000), OPENAI_PRICES["openai_sol"]) == 2.0
    assert openai_cost(
        Usage(input_tokens=1_000_000, cached_input_tokens=1_000_000),
        OPENAI_PRICES["openai_sol"],
    ) == 0.10


def test_jev_variants_have_independent_results_and_costs(tmp_path: Path) -> None:
    args = argparse.Namespace(
        number_samples=1,
        input_file=tmp_path / "input.csv",
        category_file=tmp_path / "category_input.json",
    )
    stats = {
        "typesafe_jev.recursive": Stats(successes=1, usage=Usage(input_tokens=100)),
        "typesafe_jev.recursive_subtree": Stats(wrong_predictions=1, usage=Usage(input_tokens=1000)),
    }
    output = build_output(stats, args, test_count=1)
    jev = output["models"]["typesafe_jev"]
    assert list(jev) == ["recursive", "recursive_subtree"]
    assert jev["recursive"]["success_rate"] == 1.0
    assert jev["recursive_subtree"]["success_rate"] == 0.0
    assert jev["recursive_subtree"]["strategy"] == "recursive_subtree"
    assert jev["recursive_subtree"]["model"] == jev["recursive"]["model"]
    assert jev["recursive_subtree"]["cost"]["known_total_usd"] == (
        10 * jev["recursive"]["cost"]["known_total_usd"]
    )


def test_decisions_variants_have_separate_input_only_costs(tmp_path: Path) -> None:
    import benchmark

    usage = Usage(
        input_tokens=1_000_000, cached_input_tokens=100_000,
        cache_write_tokens=200_000, output_tokens=1_000_000,
    )
    assert benchmark.openai_decisions_cost(usage) == 0.10
    args = argparse.Namespace(
        number_samples=1,
        input_file=tmp_path / "input.csv",
        category_file=tmp_path / "category_input.json",
    )
    stats = {
        "openai_luna.direct": Stats(successes=1, usage=usage),
        "openai_luna_decisions.recursive": Stats(successes=1, usage=usage),
        "openai_luna_decisions.recursive_subtree": Stats(api_errors=1),
    }
    output = build_output(stats, args, test_count=1)
    decisions = output["models"]["openai_luna_decisions"]
    assert list(decisions) == ["recursive", "recursive_subtree"]
    for result in decisions.values():
        assert result["model"] == PROVIDERS["openai_luna_decisions"].model
        assert result["api"] == "decisions"
        assert "reasoning_effort" not in result
        assert result["cost"]["input_price_usd_per_1m_tokens"] == 0.10
        assert result["cost"]["output_price_usd_per_1m_tokens"] == 0.0
    direct = output["models"]["openai_luna"]["direct"]
    assert direct["cost"]["known_total_usd"] > decisions["recursive"]["cost"]["known_total_usd"]
    assert decisions["recursive"]["cost"]["known_total_usd"] == 0.10
