#!/usr/bin/env python3
"""Download the official COICOP 2018 Excel structure and create category_input.json."""

from __future__ import annotations

import argparse
import json
import re
from collections.abc import Iterable
from pathlib import Path

import httpx
from openpyxl import load_workbook

COICOP_XLSX_URL = (
    "https://unstats.un.org/unsd/classifications/Econ/Download/"
    "COICOP_2018_English_structure.xlsx"
)
CODE_RE = re.compile(r"^\d{2}(?:\.\d+){0,5}$")
DURABILITY_SUFFIX_RE = re.compile(r"\s*\((?:ND|SD|D|S)\)\s*$", re.IGNORECASE)
HOUSEHOLD_DIVISION_MAX = 13


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\n", " ").split()).strip()


def parent_code(code: str) -> str | None:
    if "." not in code:
        return None
    return code.rsplit(".", 1)[0]


def title_key(title: str) -> str:
    """Normalize a title for semantic pass-through equality checks.

    COICOP appends durability markers such as ``(ND)``, ``(SD)``, ``(D)`` and
    ``(S)`` at some levels while an otherwise identical child can omit them.
    Those markers describe the type of expenditure rather than a different
    category meaning, so ignore a trailing marker when comparing titles.
    """
    without_marker = DURABILITY_SUFFIX_RE.sub("", clean_text(title))
    return re.sub(r"[^a-z0-9]+", " ", without_marker.lower()).strip()


def is_household_category(code: str) -> bool:
    """Return True for COICOP household expenditure divisions 01 through 13."""
    try:
        return 1 <= int(code.split(".", 1)[0]) <= HOUSEHOLD_DIVISION_MAX
    except ValueError:
        return False


def collapse_semantic_passthroughs(
    categories: list[dict[str, object]],
) -> list[dict[str, object]]:
    """Collapse single-child nodes that repeat the same semantic category.

    The official workbook can repeat the same title across adjacent levels,
    sometimes with a trailing zero code and sometimes with only a durability
    marker difference. Exposing both nodes gives classifiers duplicate valid
    answers. A node is therefore collapsed when it has exactly one *direct*
    child and the normalized titles are equivalent.

    Collapsing is transitive: for ``A -> B -> C`` where all three titles are
    equivalent, only ``C`` remains and records both ``A`` and ``B`` in
    ``collapsed_codes``. The deepest surviving category is canonical.
    """
    by_code = {str(item["code"]): item for item in categories}
    direct_children: dict[str, list[str]] = {code: [] for code in by_code}
    for code in by_code:
        pcode = parent_code(code)
        if pcode in direct_children:
            direct_children[pcode].append(code)

    # Mark all semantic pass-through parents using the original direct tree.
    redirect: dict[str, str] = {}
    for code, item in by_code.items():
        children = direct_children.get(code, [])
        if len(children) != 1:
            continue
        child_code = children[0]
        child = by_code[child_code]
        if title_key(str(item["title"])) == title_key(str(child["title"])):
            redirect[code] = child_code

    def final_target(code: str) -> str:
        seen: set[str] = set()
        while code in redirect:
            if code in seen:
                raise ValueError(f"Cycle while collapsing COICOP category {code}.")
            seen.add(code)
            code = redirect[code]
        return code

    aliases_by_target: dict[str, list[str]] = {}
    for removed_code in redirect:
        aliases_by_target.setdefault(final_target(removed_code), []).append(removed_code)

    collapsed: list[dict[str, object]] = []
    for item in categories:
        code = str(item["code"] )
        if code in redirect:
            continue
        new_item: dict[str, object] = {"code": code, "title": str(item["title"])}
        aliases = aliases_by_target.get(code)
        if aliases:
            new_item["collapsed_codes"] = sorted(aliases, key=lambda value: (value.count("."), value))
        collapsed.append(new_item)

    for idx, item in enumerate(collapsed):
        item["id"] = idx
    return collapsed


def add_tree_metadata(categories: list[dict[str, object]]) -> list[dict[str, object]]:
    """Add explicit tree fields while preserving sequential numeric category ids."""
    code_to_id = {str(item["code"]): int(item["id"]) for item in categories}
    children_by_id: dict[int, list[int]] = {int(item["id"]): [] for item in categories}

    for item in categories:
        code = str(item["code"])
        pcode = parent_code(code)
        # A collapsed `.0` node can skip its original direct parent. Walk upward
        # until the nearest surviving ancestor is found.
        while pcode is not None and pcode not in code_to_id:
            pcode = parent_code(pcode)
        pid = code_to_id.get(pcode) if pcode is not None else None
        item["parent_id"] = pid
        item["level"] = code.count(".") + 1
        item["is_optional_detail"] = item["level"] > 4
        if pid is not None:
            children_by_id[pid].append(int(item["id"]))

    for item in categories:
        cid = int(item["id"])
        children = children_by_id[cid]
        item["children_ids"] = children
        item["is_leaf"] = len(children) == 0

    return categories


def extract_categories(
    excel_path: Path,
    *,
    household_only: bool = True,
    sheet_number: int = 1,
    code_column: int = 1,
    title_column: int = 2,
) -> list[dict[str, object]]:
    """Read a header in row 1 and category rows from row 2; all indexes are 1-based.

    Defaults match the local XLSX: first worksheet, codes in A, titles in B.
    Other columns are ignored. Worksheet and column positions are not inferred.
    """
    if min(sheet_number, code_column, title_column) < 1:
        raise ValueError("Sheet number and column indexes must be at least 1.")
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    if sheet_number > len(workbook.worksheets):
        workbook.close()
        raise ValueError(f"Worksheet {sheet_number} does not exist in {excel_path}.")
    sheet = workbook.worksheets[sheet_number - 1]
    categories: list[dict[str, object]] = []
    seen_codes: set[str] = set()
    for row in sheet.iter_rows(min_row=2, max_col=max(code_column, title_column), values_only=True):
        code = clean_text(row[code_column - 1])
        title = clean_text(row[title_column - 1])
        if not CODE_RE.fullmatch(code) or not title or code in seen_codes:
            continue
        seen_codes.add(code)
        if household_only and not is_household_category(code):
            continue
        categories.append({"id": len(categories), "code": code, "title": title})
    workbook.close()

    if not categories:
        raise ValueError("No COICOP categories were extracted from the workbook.")
    categories = collapse_semantic_passthroughs(categories)
    return add_tree_metadata(categories)


def download_excel(url: str, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    with httpx.Client(follow_redirects=True, timeout=60.0) as client:
        response = client.get(url)
        response.raise_for_status()
        destination.write_bytes(response.content)


def write_json(categories: Iterable[dict[str, object]], destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        json.dumps(list(categories), indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Download COICOP 2018 and create tree-aware category_input.json.",
        epilog=(
            "Expected XLSX layout: row 1 is a header, category rows start at row 2. "
            "By default, read the first worksheet with codes in column A and titles in column B. "
            "All sheet and column indexes are 1-based. Other columns are ignored."
        ),
    )
    parser.add_argument("--url", default=COICOP_XLSX_URL, help="COICOP Excel URL")
    parser.add_argument("--sheet-number", type=int, default=1, help="Worksheet number, 1-based (default: 1)")
    parser.add_argument("--code-column", type=int, default=1, help="Code column index, 1-based (default: 1 = A)")
    parser.add_argument("--title-column", type=int, default=2, help="Title column index, 1-based (default: 2 = B)")
    parser.add_argument(
        "--excel-output",
        type=Path,
        default=Path("COICOP_2018_English_structure.xlsx"),
        help="Where to save the downloaded Excel file",
    )
    parser.add_argument(
        "--json-output",
        type=Path,
        default=Path("category_input.json"),
        help="Where to write the generated category JSON",
    )
    parser.add_argument(
        "--reuse-excel",
        action="store_true",
        help="Do not download when --excel-output already exists",
    )
    parser.add_argument(
        "--include-all-divisions",
        action="store_true",
        help=(
            "Include COICOP divisions 14-15 as well. By default only household "
            "expenditure divisions 01-13 are emitted, which is appropriate for "
            "consumer expense classification."
        ),
    )
    args = parser.parse_args()
    if min(args.sheet_number, args.code_column, args.title_column) < 1:
        parser.error("--sheet-number, --code-column and --title-column must be at least 1.")
    return args


def main() -> None:
    args = parse_args()
    if not (args.reuse_excel and args.excel_output.exists()):
        print(f"Downloading {args.url}")
        download_excel(args.url, args.excel_output)
    categories = extract_categories(
        args.excel_output,
        household_only=not args.include_all_divisions,
        sheet_number=args.sheet_number,
        code_column=args.code_column,
        title_column=args.title_column,
    )
    write_json(categories, args.json_output)
    scope = "all COICOP divisions" if args.include_all_divisions else "household divisions 01-13"
    print(
        f"Wrote {len(categories)} collapsed, tree-aware categories ({scope}) "
        f"to {args.json_output}"
    )


if __name__ == "__main__":
    main()
